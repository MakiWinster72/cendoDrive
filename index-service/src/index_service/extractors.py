"""文本提取。

对应"可先假设/先做"清单第 2、3、9 项与对接说明 第 7 节：

- 支持 PDF、DOCX、TXT、Markdown；其余格式抛 ``UnsupportedFormat``；
- TXT / Markdown 的编码兜底：UTF-8(BOM) → GB18030 → UTF-8 替换，
  并统一 CRLF 换行为 LF（实现补充，见 README 假设清单）；
- 单文件大小上限 50MB、PDF 页数上限 1000、单文件解析超时 120s（均可配置）；
- 空文本（含扫描件 PDF 无文字层）抛 ``EmptyText``。

异常分类决定后续是重试还是终止，见 第 6 节：格式损坏与超限不重试。
"""

from __future__ import annotations

import time
from dataclasses import dataclass
from pathlib import Path

from .config import Settings

SUPPORTED_EXTENSIONS = frozenset({"pdf", "docx", "txt", "md", "markdown"})

# 编码探测顺序：UTF-8 含 BOM → GB18030（兼容 GBK/GB2312）
_ENCODINGS = ("utf-8-sig", "gb18030")


class ExtractionError(Exception):
    """提取失败的基类。"""


class UnsupportedFormat(ExtractionError):
    """扩展名不在支持范围内。"""


class CorruptDocument(ExtractionError):
    """文件损坏或无法解析。不重试。"""


class ParseTimeout(ExtractionError):
    """解析超出时间预算。不重试。"""


class TooLargeFile(ExtractionError):
    """超过单文件大小上限。不重试。"""


class TooManyPages(ExtractionError):
    """PDF 页数超过上限。不重试。"""


class EmptyText(ExtractionError):
    """没有可用文本（空文件或扫描件）。"""


@dataclass(frozen=True, slots=True)
class ExtractionResult:
    text: str
    pages: int | None
    encoding: str | None


def extension_of(name: str) -> str:
    _, dot, suffix = name.rpartition(".")
    return suffix.lower() if dot else ""


def decode_text(raw: bytes) -> tuple[str, str]:
    """按 UTF-8(BOM) → GB18030 顺序解码，全部失败时以替换字符兜底。"""
    for encoding in _ENCODINGS:
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8/replace"


def normalize_newlines(text: str) -> str:
    """统一换行符，避免 Windows CRLF 的 ``\\r`` 残留到分块文本与搜索片段中。"""
    if "\r" not in text:
        return text
    return text.replace("\r\n", "\n").replace("\r", "\n")


def extract(path: Path, name: str, settings: Settings) -> ExtractionResult:
    """按扩展名分发到具体提取器。"""
    extension = extension_of(name)
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormat(f"不支持的文档类型: .{extension or '(无扩展名)'}")

    size = path.stat().st_size
    if size > settings.max_file_bytes:
        raise TooLargeFile(f"文件大小 {size} 字节，超过上限 {settings.max_file_bytes} 字节")

    if extension == "pdf":
        return _extract_pdf(path, settings)
    if extension == "docx":
        return _extract_docx(path)
    return _extract_plain(path)


def _extract_plain(path: Path) -> ExtractionResult:
    text, encoding = decode_text(path.read_bytes())
    if not text.strip():
        raise EmptyText("TXT/Markdown 没有可用文本")
    return ExtractionResult(text=normalize_newlines(text), pages=None, encoding=encoding)


def _extract_pdf(path: Path, settings: Settings) -> ExtractionResult:
    import pdfplumber

    deadline = time.monotonic() + settings.parse_timeout_seconds
    try:
        with pdfplumber.open(path) as pdf:
            page_count = len(pdf.pages)
            if page_count > settings.max_pdf_pages:
                raise TooManyPages(f"PDF 页数 {page_count} 超过上限 {settings.max_pdf_pages}")
            parts: list[str] = []
            for page in pdf.pages:
                if time.monotonic() > deadline:
                    raise ParseTimeout("PDF 解析超出时间预算")
                parts.append(page.extract_text() or "")
    except (TooManyPages, ParseTimeout):
        raise
    except Exception as exc:  # pdfminer 会抛出多种解析异常
        raise CorruptDocument(f"PDF 解析失败: {exc}") from exc

    text = "\n".join(parts)
    if not text.strip():
        raise EmptyText("PDF 没有文字层（可能是扫描件）")
    return ExtractionResult(text=text, pages=page_count, encoding=None)


def _extract_docx(path: Path) -> ExtractionResult:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise CorruptDocument(f"DOCX 解析失败: {exc}") from exc

    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))

    text = "\n".join(parts)
    if not text.strip():
        raise EmptyText("DOCX 没有可用文本")
    return ExtractionResult(text=text, pages=None, encoding=None)
