"""测试辅助：样例文件构造与事件构造。

PDF 样例在测试时按字节生成，避免把二进制文件提交进仓库。
"""

from __future__ import annotations

from pathlib import Path

from index_service.models import IndexEvent, Operation


def _escape_pdf_text(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def minimal_pdf(pages: list[str]) -> bytes:
    """生成结构合法的最小 PDF（内置 Helvetica，ASCII 文本），用于覆盖 PDF 解析路径。"""
    font_id = 3
    page_ids: list[int] = []
    content_ids: list[int] = []
    next_id = 4
    for _ in pages:
        page_ids.append(next_id)
        next_id += 1
        content_ids.append(next_id)
        next_id += 1

    objects: dict[int, bytes] = {}
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[2] = f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode("ascii")
    objects[font_id] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    for index, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 20 150 Td ({_escape_pdf_text(text)}) Tj ET".encode("latin-1")
        objects[content_ids[index]] = (
            b"<< /Length "
            + str(len(stream)).encode("ascii")
            + b" >>\nstream\n"
            + stream
            + b"\nendstream"
        )
        objects[page_ids[index]] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 300] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
            f"/Contents {content_ids[index]} 0 R >>"
        ).encode("ascii")

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for object_id in sorted(objects):
        offsets[object_id] = len(out)
        out += f"{object_id} 0 obj\n".encode("ascii") + objects[object_id] + b"\nendobj\n"

    xref_offset = len(out)
    size = max(objects) + 1
    out += f"xref\n0 {size}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for object_id in range(1, size):
        out += f"{offsets[object_id]:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n"
    ).encode("ascii")
    return bytes(out)


def write_pdf(path: Path, pages: list[str]) -> Path:
    path.write_bytes(minimal_pdf(pages))
    return path


def write_docx(path: Path, paragraphs: list[str], table: list[list[str]] | None = None) -> Path:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        created = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for cell_index, value in enumerate(row):
                created.cell(row_index, cell_index).text = value
    document.save(str(path))
    return path


def make_upsert(
    *,
    file_id: str = "1",
    owner_id: str = "1",
    revision: int = 1,
    name: str = "notes.txt",
    key: str | None = None,
    backend: str = "local",
    event_id: str | None = None,
) -> IndexEvent:
    return IndexEvent(
        eventId=event_id or f"file-{file_id}-{revision}-UPSERT",
        operation=Operation.UPSERT,
        fileId=file_id,
        ownerId=owner_id,
        revision=revision,
        fileName=name,
        storageBackend=backend,
        storageKey=key if key is not None else name,
    )


def make_lifecycle(
    operation: str | Operation,
    *,
    file_id: str = "1",
    owner_id: str = "1",
    revision: int = 2,
    event_id: str | None = None,
) -> IndexEvent:
    resolved = Operation(operation)
    return IndexEvent(
        eventId=event_id or f"file-{file_id}-{revision}-{resolved.value}",
        operation=resolved,
        fileId=file_id,
        ownerId=owner_id,
        revision=revision,
    )
